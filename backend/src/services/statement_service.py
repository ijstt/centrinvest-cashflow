import csv
import io
import json
import hashlib
import openpyxl
import re
from datetime import datetime, date
from decimal import Decimal
from typing import Literal
from sqlalchemy import select, and_, func, desc
from sqlalchemy.ext.asyncio import AsyncSession
from src.models.entities import (
    Company, TransactionRecord, DailyActualRecord, CategoryOverride, ImportRecord
)
from src.schemas.contracts import (
    ColumnMapping, ImportResponse, ImportIssue, Transaction,
    SourceReference, TransactionListResponse, ActualsResponse, DailyActual, CategoryTotal
)
from src.services.category_service import CategoryService

def sanitize_html_tags(val: str | None) -> str | None:
    if val is None:
        return None
    s = str(val)
    # Strip script blocks
    s = re.sub(r"(?is)<script\b[^>]*>.*?</script>", "", s)
    # Strip standalone script tags
    s = re.sub(r"(?is)</?script\b[^>]*>", "", s)
    # Strip inline script / event handlers
    s = re.sub(r"(?i)[/\s]+on\w+\s*=\s*(?:\"[^\"]*\"|\x27[^\x27]*\x27|[^\s>]+)", "", s)
    return s.strip()

def neutralize_formula(val: any) -> str:
    if val is None:
        return ""
    s = str(val)
    if s and s[0] in ("=", "@", "+", "-"):
        return f"'{s}"
    return s

def decode_text(raw_bytes: bytes) -> str:
    try:
        return raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return raw_bytes.decode("windows-1251")
    except UnicodeDecodeError:
        pass
    return raw_bytes.decode("utf-8-sig", errors="replace")

def parse_date_object(val: str | date | datetime) -> date:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()
    if len(s) >= 10:
        try:
            return date.fromisoformat(s[:10])
        except ValueError:
            pass
    return datetime.fromisoformat(s.replace("Z", "+00:00")).date()

def is_1c_file(text: str) -> bool:
    return text.strip().startswith("1CClientBankExchange")

def parse_1c_exchange(text: str, default_closing_balance: float = 0.0) -> tuple[list[dict], ColumnMapping]:
    rows = []
    lines = text.splitlines()
    in_doc = False
    initial_bal = None
    closing_bal = None
    header_accounts = set()
    header_inns = set()

    for line in lines:
        line = line.strip()
        if line.startswith("НачальныйОстаток="):
            val = parse_decimal(line.split("=", 1)[1])
            if val is not None:
                initial_bal = val
        elif line.startswith("КонечныйОстаток="):
            val = parse_decimal(line.split("=", 1)[1])
            if val is not None:
                closing_bal = val
        elif line.startswith("РасчСчет="):
            acc = line.split("=", 1)[1].strip()
            if acc:
                header_accounts.add(acc)
        elif line.startswith("ИНН="):
            inn = line.split("=", 1)[1].strip()
            if inn:
                header_inns.add(inn)

    raw_docs = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith("СекцияДокумент="):
            in_doc = True
            current_doc = {"doc_type": line.split("=", 1)[1].strip()}
        elif line == "КонецДокумента":
            if in_doc and current_doc:
                d_val = parse_date(current_doc.get("Дата"))
                amt = parse_decimal(current_doc.get("Сумма"))
                payer_acc = current_doc.get("ПлательщикСчет") or current_doc.get("ПлательщикРасчСчет") or ""
                receiver_acc = current_doc.get("ПолучательСчет") or current_doc.get("ПолучательРасчСчет") or ""
                payer = current_doc.get("Плательщик") or current_doc.get("Плательщик1") or ""
                payer_inn = current_doc.get("ПлательщикИНН") or ""
                receiver = current_doc.get("Получатель") or current_doc.get("Получатель1") or ""
                receiver_inn = current_doc.get("ПолучательИНН") or ""
                purpose = current_doc.get("НазначениеПлатежа") or ""
                doc_num = current_doc.get("Номер") or ""
                ext_id = f"{d_val}_{doc_num}" if (d_val and doc_num) else (doc_num or None)

                is_income = False
                if header_accounts:
                    if receiver_acc in header_accounts and payer_acc not in header_accounts:
                        is_income = True
                    elif payer_acc in header_accounts and receiver_acc not in header_accounts:
                        is_income = False
                    elif header_inns:
                        if receiver_inn in header_inns and payer_inn not in header_inns:
                            is_income = True
                        elif payer_inn in header_inns:
                            is_income = False
                elif header_inns:
                    if receiver_inn in header_inns and payer_inn not in header_inns:
                        is_income = True
                    elif payer_inn in header_inns:
                        is_income = False

                direction = "income" if is_income else "expense"
                c_name = (payer or receiver) if is_income else (receiver or payer)
                c_inn = (payer_inn or receiver_inn) if is_income else (receiver_inn or payer_inn)

                raw_docs.append({
                    "date": d_val,
                    "amount": amt,
                    "direction": direction,
                    "counterparty_name": c_name,
                    "counterparty_inn": c_inn,
                    "purpose": purpose,
                    "external_id": ext_id
                })
                current_doc = {}
            in_doc = False
        elif in_doc:
            if "=" in line:
                k, v = line.split("=", 1)
                current_doc[k.strip()] = v.strip()

    # Determine starting balance for rolling calculation
    if initial_bal is not None:
        running_bal = initial_bal
    elif closing_bal is not None:
        total_delta = sum(
            (d["amount"] or 0.0) if d["direction"] == "income" else -(d["amount"] or 0.0)
            for d in raw_docs
        )
        running_bal = round(closing_bal - total_delta, 2)
    else:
        running_bal = default_closing_balance

    # Calculate true daily closing balances by rolling sequentially through transactions.
    # Group preserving order of transactions per date:
    date_docs: dict[str | None, list[dict]] = {}
    for d in raw_docs:
        dt = d["date"]
        if dt not in date_docs:
            date_docs[dt] = []
        date_docs[dt].append(d)

    # Sort dates chronologically so rolling balance is strictly time-ordered
    sorted_dates = sorted([dt for dt in date_docs.keys() if dt is not None])
    if None in date_docs:
        sorted_dates.append(None)

    for dt in sorted_dates:
        for d in date_docs[dt]:
            amt_val = d["amount"] or 0.0
            if d["direction"] == "income":
                running_bal += amt_val
            else:
                running_bal -= amt_val
        daily_close = round(running_bal, 2)
        for d in date_docs[dt]:
            d["closing_balance"] = daily_close
            rows.append(d)

    mapping = ColumnMapping(
        date="date",
        amount="amount",
        direction="direction",
        counterparty_name="counterparty_name",
        counterparty_inn="counterparty_inn",
        purpose="purpose",
        closing_balance="closing_balance",
        external_id="external_id"
    )
    return rows, mapping

# Canonical alias as specified in audit requirements
parse_1c_client_bank_exchange = parse_1c_exchange


def parse_direction(val: str) -> Literal["income", "expense"] | None:
    v = str(val).strip().lower()
    if v in ("поступление", "income", "приход", "+", "credit"):
        return "income"
    if v in ("списание", "expense", "расход", "-", "debit"):
        return "expense"
    return None

def parse_decimal(val) -> float | None:
    if val is None:
        return None
    s = str(val).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        f = float(s)
        if not (f == f and f != float("inf") and f != float("-inf")):
            return None
        return round(f, 2)
    except ValueError:
        return None

def parse_date(val) -> str | None:
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if len(s) >= 10:
        try:
            return date.fromisoformat(s[:10]).strftime("%Y-%m-%d")
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).strftime("%Y-%m-%d")
    except ValueError:
        return None

class StatementService:
    @staticmethod
    async def import_file(
        db: AsyncSession,
        company_id: str,
        file_bytes: bytes,
        format_type: Literal["csv", "xlsx", "1c"],
        mapping: ColumnMapping,
        balance_policy: Literal["reject_conflicts", "use_last_row"] = "reject_conflicts"
    ) -> ImportResponse:
        issues: list[ImportIssue] = []
        raw_rows: list[dict] = []

        if len(file_bytes) > 10 * 1024 * 1024:
            return ImportResponse(
                id="imp_oversize",
                company_id=company_id,
                status="rejected",
                accepted_rows=0,
                duplicate_rows=0,
                rejected_rows=0,
                period_start=None,
                period_end=None,
                issues=[ImportIssue(row=None, field=None, code="too_large", message="File exceeds 10 MiB limit")],
                data_revision=0
            )

        # Serialize the entire company import before checking existing operation IDs.
        comp = (await db.execute(select(Company).where(
            Company.id == company_id).with_for_update())).scalar_one_or_none()
        if not comp:
            raise KeyError(f'Company {company_id} not found')
        try:
            if format_type in ("csv", "1c"):
                text = decode_text(file_bytes)
                if is_1c_file(text) or format_type == "1c":
                    comp_q = await db.execute(select(Company).where(Company.id == company_id))
                    comp_obj = comp_q.scalar_one_or_none()
                    fallback_bal = float(comp_obj.closing_balance) if (comp_obj and comp_obj.closing_balance) else 0.0
                    raw_rows, mapping = parse_1c_exchange(text, default_closing_balance=fallback_bal)
                else:
                    header_line = text.splitlines()[0] if text.splitlines() else ""
                    semis = header_line.count(";")
                    commas = header_line.count(",")
                    tabs = header_line.count("\t")
                    if tabs > semis and tabs > commas:
                        delimiter = "\t"
                    elif semis > commas:
                        delimiter = ";"
                    else:
                        delimiter = ","
                    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
                    for r in reader:
                        raw_rows.append(r)
            else:
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                ws = wb.active
                header = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    row_dict = {}
                    for col_name, val in zip(header, row):
                        if col_name:
                            row_dict[col_name] = val
                    raw_rows.append(row_dict)
        except Exception as e:
            return ImportResponse(
                id="imp_corrupt",
                company_id=company_id,
                status="rejected",
                accepted_rows=0,
                duplicate_rows=0,
                rejected_rows=0,
                issues=[ImportIssue(row=None, field=None, code="corrupt_file", message=str(e))],
                data_revision=0
            )

        if len(raw_rows) > 100_000:
            return ImportResponse(
                id="imp_too_many_rows",
                company_id=company_id,
                status="rejected",
                accepted_rows=0,
                duplicate_rows=0,
                rejected_rows=len(raw_rows),
                issues=[ImportIssue(row=None, field=None, code="too_large", message="Row count exceeds 100,000")],
                data_revision=0
            )

        if not raw_rows:
            return ImportResponse(
                id="imp_empty",
                company_id=company_id,
                status="rejected",
                accepted_rows=0,
                duplicate_rows=0,
                rejected_rows=0,
                issues=[ImportIssue(row=None, field=None, code="empty_file", message="Statement has zero rows")],
                data_revision=0
            )

        # Parse and validate rows
        parsed_records = []
        daily_closing_balances: dict[str, list[float]] = {}
        content_sig_counts: dict[str, int] = {}

        for idx, row in enumerate(raw_rows, start=1):
            date_val = parse_date(row.get(mapping.date))
            if not date_val:
                issues.append(ImportIssue(row=idx, field=mapping.date, code="invalid_date", message="Invalid or missing date"))
                continue

            raw_amt = parse_decimal(row.get(mapping.amount))
            if raw_amt is None or raw_amt == 0:
                issues.append(ImportIssue(row=idx, field=mapping.amount, code="invalid_amount", message="Amount must be non-zero number"))
                continue

            amt = abs(raw_amt)
            direction = parse_direction(row.get(mapping.direction))
            if not direction and raw_amt < 0:
                direction = "expense"
            elif not direction and raw_amt > 0:
                direction = "income"

            if not direction:
                issues.append(ImportIssue(row=idx, field=mapping.direction, code="invalid_direction", message="Direction must be income/expense"))
                continue

            c_name = sanitize_html_tags(str(row.get(mapping.counterparty_name) or "").strip()) or ""
            raw_inn = row.get(mapping.counterparty_inn)
            if raw_inn is not None:
                if isinstance(raw_inn, float) and raw_inn.is_integer():
                    c_inn = str(int(raw_inn)).zfill(10)
                elif isinstance(raw_inn, int):
                    c_inn = str(raw_inn).zfill(10)
                else:
                    c_inn = str(raw_inn).strip()
                    if c_inn.isdigit() and len(c_inn) <= 9:
                        c_inn = c_inn.zfill(10)
            else:
                c_inn = ""
            purpose = sanitize_html_tags(str(row.get(mapping.purpose) or "").strip()) or ""

            bal = parse_decimal(row.get(mapping.closing_balance))
            if bal is None:
                issues.append(ImportIssue(row=idx, field=mapping.closing_balance, code="invalid_balance", message="Invalid closing balance"))
                continue

            ext_id = None
            if mapping.external_id and mapping.external_id in row:
                ext_id = str(row.get(mapping.external_id) or "").strip() or None

            # If external_id is missing, compute deterministic content hash with occurrence index
            # to preserve identical valid transactions (e.g. retail checks) within the same file,
            # while remaining strictly idempotent across re-imports of the file.
            if not ext_id:
                content_sig = f"{date_val}|{amt:.2f}|{direction}|{c_name}|{c_inn}|{purpose[:64]}"
                occ = content_sig_counts.get(content_sig, 0)
                content_sig_counts[content_sig] = occ + 1
                ext_id = f"hash_{hashlib.sha256(f'{content_sig}|occ={occ}'.encode('utf-8')).hexdigest()[:24]}"

            daily_closing_balances.setdefault(date_val, []).append(bal)

            parsed_records.append({
                "row_idx": idx,
                "date": date_val,
                "amount": amt,
                "direction": direction,
                "counterparty_name": c_name,
                "counterparty_inn": c_inn,
                "purpose": purpose,
                "closing_balance": bal,
                "external_id": ext_id
            })

        # Check closing balance consistency per day
        for d, bals in daily_closing_balances.items():
            first_bal = bals[0]
            conflicts = [b for b in bals if abs(b - first_bal) > 0.001]
            if conflicts:
                if balance_policy == "reject_conflicts":
                    issues.append(ImportIssue(
                        row=None,
                        field=mapping.closing_balance,
                        code="conflicting_daily_balance",
                        message=f"Conflicting closing balances on date {d}: {bals}"
                    ))
                else:
                    issues.append(ImportIssue(
                        row=None,
                        field=mapping.closing_balance,
                        code="balance_warning",
                        message=f"Multiple closing balances on {d}; using last row balance: {bals[-1]}"
                    ))

        # Check if rejected due to critical issues or row validation errors (enforce atomic import invariant)
        critical_issues = [
            iss for iss in issues
            if iss.code in (
                "conflicting_daily_balance", "empty_file", "too_large",
                "invalid_date", "invalid_amount", "invalid_direction", "invalid_balance"
            )
        ]
        if critical_issues:
            return ImportResponse(
                id="imp_rejected",
                company_id=company_id,
                status="rejected",
                accepted_rows=0,
                duplicate_rows=0,
                rejected_rows=len(raw_rows),
                issues=issues,
                data_revision=0
            )

        # Load manual category overrides for company
        override_res = await db.execute(
            select(CategoryOverride).where(CategoryOverride.company_id == company_id)
        )
        overrides = {ov.external_id: ov.category_id for ov in override_res.scalars().all()}

        # Load existing external_ids for company to deduplicate
        existing_ext_res = await db.execute(
            select(TransactionRecord.external_id).where(
                TransactionRecord.company_id == company_id,
                TransactionRecord.external_id.isnot(None)
            )
        )
        existing_ext_ids = {r[0] for r in existing_ext_res.all()}

        import_rec = ImportRecord(
            company_id=company_id,
            status="ready",
            issues=[iss.model_dump() for iss in issues]
        )
        db.add(import_rec)
        await db.flush()

        accepted_count = 0
        duplicate_count = 0
        new_tx_records = []
        all_dates = []

        for p in parsed_records:
            if p["external_id"] and p["external_id"] in existing_ext_ids:
                duplicate_count += 1
                continue

            all_dates.append(p["date"])
            if p["external_id"]:
                existing_ext_ids.add(p["external_id"])

            # Determine category
            if p["external_id"] and p["external_id"] in overrides:
                cat_id = overrides[p["external_id"]]
                cat_method = "manual"
                cat_conf = 1.0
            else:
                cat_id, cat_method, cat_conf = CategoryService.categorize(
                    direction=p["direction"],
                    purpose=p["purpose"],
                    counterparty_name=p["counterparty_name"],
                    counterparty_inn=p["counterparty_inn"]
                )

            tx = TransactionRecord(
                company_id=company_id,
                import_id=import_rec.id,
                date=datetime.strptime(p["date"], "%Y-%m-%d").date(),
                amount=p["amount"],
                direction=p["direction"],
                counterparty_name=sanitize_html_tags(p["counterparty_name"]) or "",
                counterparty_inn=p["counterparty_inn"],
                purpose=sanitize_html_tags(p["purpose"]) or "",
                closing_balance=p["closing_balance"],
                external_id=p["external_id"],
                category_id=cat_id,
                category_method=cat_method,
                category_confidence=cat_conf,
                source_kind="statement",
                source_id=import_rec.id,
                source_label=f"Выписка {p['date']}"
            )
            new_tx_records.append(tx)
            accepted_count += 1

        db.add_all(new_tx_records)
        await db.flush()

        # Update daily actuals for all dates in statement
        # Compute daily totals from db for company
        comp_res = await db.execute(select(Company).where(Company.id == company_id).with_for_update())
        comp = comp_res.scalar_one()

        if accepted_count:
            comp.data_revision += 1

        # Recompute daily actuals from all transactions (excluding sandbox)
        daily_tx_res = await db.execute(
            select(
                TransactionRecord.date,
                func.sum(case_income()).label("income"),
                func.sum(case_expense()).label("expense"),
                func.count(TransactionRecord.id).label("count")
            ).where(
                TransactionRecord.company_id == company_id,
                (TransactionRecord.source_kind != "sandbox") | (TransactionRecord.source_kind.is_(None))
            ).group_by(TransactionRecord.date).order_by(TransactionRecord.date)
        )

        # Get latest closing balance per date (excluding sandbox)
        latest_bal_res = await db.execute(
            select(TransactionRecord.date, TransactionRecord.closing_balance)
            .where(
                TransactionRecord.company_id == company_id,
                (TransactionRecord.source_kind != "sandbox") | (TransactionRecord.source_kind.is_(None))
            )
            .order_by(TransactionRecord.date, desc(TransactionRecord.created_at), desc(TransactionRecord.id))
        )
        date_to_bal = {}
        for d, b in latest_bal_res.all():
            if d not in date_to_bal:
                date_to_bal[d] = float(b)

        last_date = None
        last_bal = None
        valid_dates = set()

        for d, inc, exp, cnt in daily_tx_res.all():
            valid_dates.add(d)
            last_date = d
            c_bal = date_to_bal.get(d, 0.0)
            last_bal = c_bal
            # Upsert DailyActualRecord
            act_res = await db.execute(
                select(DailyActualRecord).where(
                    DailyActualRecord.company_id == company_id,
                    DailyActualRecord.date == d
                )
            )
            actual_rec = act_res.scalar_one_or_none()
            if not actual_rec:
                actual_rec = DailyActualRecord(
                    company_id=company_id,
                    date=d,
                    income=float(inc or 0.0),
                    expense=float(exp or 0.0),
                    closing_balance=c_bal,
                    transaction_count=cnt,
                    observation="observed"
                )
                db.add(actual_rec)
            else:
                actual_rec.income = float(inc or 0.0)
                actual_rec.expense = float(exp or 0.0)
                actual_rec.closing_balance = c_bal
                actual_rec.transaction_count = cnt

        if valid_dates:
            from sqlalchemy import delete
            await db.execute(
                delete(DailyActualRecord).where(
                    DailyActualRecord.company_id == company_id,
                    DailyActualRecord.date.notin_(valid_dates)
                )
            )

        if last_date:
            comp.balance_date = last_date
            comp.closing_balance = last_bal

        period_start = min(all_dates) if all_dates else None
        period_end = max(all_dates) if all_dates else None

        import_rec.accepted_rows = accepted_count
        import_rec.duplicate_rows = duplicate_count
        import_rec.rejected_rows = len(issues)
        import_rec.period_start = datetime.strptime(period_start, "%Y-%m-%d").date() if period_start else None
        import_rec.period_end = datetime.strptime(period_end, "%Y-%m-%d").date() if period_end else None
        import_rec.data_revision = comp.data_revision

        from src.services.financial_mutation import complete_financial_change
        await complete_financial_change(db, company_id,
            min(all_dates) if all_dates else (comp.balance_date or date.today()).isoformat(),
            refresh=bool(accepted_count))

        return ImportResponse(
            id=import_rec.id,
            company_id=company_id,
            status="ready",
            accepted_rows=accepted_count,
            duplicate_rows=duplicate_count,
            rejected_rows=len(issues),
            period_start=period_start,
            period_end=period_end,
            issues=issues,
            data_revision=comp.data_revision
        )

    @staticmethod
    async def list_transactions(
        db: AsyncSession,
        company_id: str,
        limit: int = 50,
        cursor: str | None = None,
        from_date: str | None = None,
        to_date: str | None = None,
        direction: str | None = None,
        category_id: str | None = None,
        search: str | None = None
    ) -> TransactionListResponse:
        conditions = [
            TransactionRecord.company_id == company_id,
            (TransactionRecord.source_kind != "sandbox") | (TransactionRecord.source_kind.is_(None))
        ]
        if from_date:
            conditions.append(TransactionRecord.date >= parse_date_object(from_date))
        if to_date:
            conditions.append(TransactionRecord.date <= parse_date_object(to_date))
        if direction:
            conditions.append(TransactionRecord.direction == direction)
        if category_id:
            conditions.append(TransactionRecord.category_id == category_id)
        if search:
            escaped_search = search.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
            s = f"%{escaped_search}%"
            conditions.append(
                (TransactionRecord.purpose.ilike(s, escape="\\")) |
                (TransactionRecord.counterparty_name.ilike(s, escape="\\")) |
                (TransactionRecord.counterparty_inn.ilike(s, escape="\\"))
            )

        # Count total
        count_q = select(func.count(TransactionRecord.id)).where(and_(*conditions))
        total_res = await db.execute(count_q)
        total = total_res.scalar() or 0

        # Query items
        q = (
            select(TransactionRecord)
            .where(and_(*conditions))
            .order_by(desc(TransactionRecord.date), desc(TransactionRecord.id))
            .limit(limit)
        )
        if cursor:
            # Simple offset cursor or id cursor
            try:
                offset = int(cursor)
                q = q.offset(offset)
            except ValueError:
                pass

        res = await db.execute(q)
        records = res.scalars().all()

        valid_source_kinds = {
            "statement", "manual", "invoice", "contract", "tax",
            "bank", "accounting", "cash_register", "marketplace", "registry"
        }

        items = []
        for r in records:
            sk = r.source_kind if r.source_kind in valid_source_kinds else "statement"
            items.append(Transaction(
                id=r.id,
                company_id=r.company_id,
                date=r.date.strftime("%Y-%m-%d"),
                amount=float(r.amount),
                direction=r.direction,
                counterparty_name=r.counterparty_name,
                counterparty_inn=r.counterparty_inn,
                purpose=r.purpose,
                closing_balance=float(r.closing_balance),
                external_id=r.external_id,
                category_id=r.category_id,
                category_method=r.category_method,
                category_confidence=r.category_confidence,
                source=SourceReference(
                    kind=sk,
                    source_id=r.source_id,
                    external_id=r.external_id,
                    label=r.source_label or "Выписка"
                ),
                matched_event_id=r.matched_event_id
            ))

        next_cursor = None
        cur_offset = int(cursor) if cursor and cursor.isdigit() else 0
        if cur_offset + len(items) < total:
            next_cursor = str(cur_offset + limit)

        return TransactionListResponse(
            items=items,
            next_cursor=next_cursor,
            total=total
        )

    @staticmethod
    async def update_category(
        db: AsyncSession,
        company_id: str,
        transaction_id: str,
        category_id: str
    ) -> Transaction:
        res = await db.execute(
            select(TransactionRecord).where(
                TransactionRecord.company_id == company_id,
                TransactionRecord.id == transaction_id
            )
        )
        tx = res.scalar_one_or_none()
        if not tx:
            raise KeyError(f"Transaction {transaction_id} not found in company {company_id}")

        tx.category_id = category_id
        tx.category_method = "manual"
        tx.category_confidence = 1.0

        if tx.external_id:
            # Save or update CategoryOverride
            ov_res = await db.execute(
                select(CategoryOverride).where(
                    CategoryOverride.company_id == company_id,
                    CategoryOverride.external_id == tx.external_id
                )
            )
            ov = ov_res.scalar_one_or_none()
            if not ov:
                ov = CategoryOverride(
                    company_id=company_id,
                    external_id=tx.external_id,
                    category_id=category_id
                )
                db.add(ov)
            else:
                ov.category_id = category_id

        comp_res = await db.execute(select(Company).where(Company.id == company_id))
        comp = comp_res.scalar_one()
        comp.data_revision += 1

        await db.commit()
        await db.refresh(tx)

        return Transaction(
            id=tx.id,
            company_id=tx.company_id,
            date=tx.date.strftime("%Y-%m-%d"),
            amount=float(tx.amount),
            direction=tx.direction,
            counterparty_name=tx.counterparty_name,
            counterparty_inn=tx.counterparty_inn,
            purpose=tx.purpose,
            closing_balance=float(tx.closing_balance),
            external_id=tx.external_id,
            category_id=tx.category_id,
            category_method=tx.category_method,
            category_confidence=tx.category_confidence,
            source=SourceReference(
                kind="statement",
                source_id=tx.source_id,
                external_id=tx.external_id,
                label=tx.source_label or "Выписка"
            ),
            matched_event_id=tx.matched_event_id
        )

    @staticmethod
    async def get_actuals(
        db: AsyncSession,
        company_id: str,
        from_date: str | None = None,
        to_date: str | None = None,
        granularity: str = "day"
    ) -> ActualsResponse:
        comp_res = await db.execute(select(Company).where(Company.id == company_id))
        comp = comp_res.scalar_one()

        conditions = [DailyActualRecord.company_id == company_id]
        if from_date:
            conditions.append(DailyActualRecord.date >= parse_date_object(from_date))
        if to_date:
            conditions.append(DailyActualRecord.date <= parse_date_object(to_date))

        res = await db.execute(
            select(DailyActualRecord).where(and_(*conditions)).order_by(DailyActualRecord.date)
        )
        days = []
        for r in res.scalars().all():
            days.append(DailyActual(
                date=r.date.strftime("%Y-%m-%d"),
                income=float(r.income) if r.income is not None else None,
                expense=float(r.expense) if r.expense is not None else None,
                closing_balance=float(r.closing_balance) if r.closing_balance is not None else None,
                transaction_count=r.transaction_count,
                observation=r.observation
            ))

        # Calculate category totals
        cat_conditions = [
            TransactionRecord.company_id == company_id,
            (TransactionRecord.source_kind != "sandbox") | (TransactionRecord.source_kind.is_(None))
        ]
        if from_date:
            cat_conditions.append(TransactionRecord.date >= parse_date_object(from_date))
        if to_date:
            cat_conditions.append(TransactionRecord.date <= parse_date_object(to_date))

        cat_res = await db.execute(
            select(
                TransactionRecord.category_id,
                TransactionRecord.direction,
                func.sum(TransactionRecord.amount).label("amount"),
                func.count(TransactionRecord.id).label("count")
            ).where(and_(*cat_conditions)).group_by(
                TransactionRecord.category_id, TransactionRecord.direction
            )
        )

        category_totals = []
        for c_id, dirn, amt, cnt in cat_res.all():
            category_totals.append(CategoryTotal(
                category_id=c_id,
                direction=dirn,
                amount=round(float(amt or 0.0), 2),
                count=cnt
            ))

        return ActualsResponse(
            company_id=company_id,
            data_revision=comp.data_revision,
            days=days,
            category_totals=category_totals
        )

    @staticmethod
    def neutralize_formula_cell(val: any) -> str:
        return neutralize_formula(val)

    @staticmethod
    async def export_csv(
        db: AsyncSession,
        company_id: str,
        from_date: str | None = None,
        to_date: str | None = None
    ) -> str:
        conditions = [TransactionRecord.company_id == company_id]
        if from_date:
            parsed_from = parse_date_object(from_date)
            if parsed_from:
                conditions.append(TransactionRecord.date >= parsed_from)
        if to_date:
            parsed_to = parse_date_object(to_date)
            if parsed_to:
                conditions.append(TransactionRecord.date <= parsed_to)

        res = await db.execute(
            select(TransactionRecord)
            .where(and_(*conditions))
            .order_by(TransactionRecord.date.desc(), TransactionRecord.id)
        )
        transactions = res.scalars().all()
        return StatementService.export_transactions_csv(transactions)

    @staticmethod
    def export_transactions_csv(transactions: list) -> str:
        output = io.StringIO()
        writer = csv.writer(output, delimiter=";")
        writer.writerow([
            "date", "amount", "direction", "counterparty_name",
            "counterparty_inn", "purpose", "closing_balance", "category_id"
        ])
        for tx in transactions:
            if isinstance(tx, dict):
                d = tx.get("date", "")
                amt = tx.get("amount", "")
                dirn = tx.get("direction", "")
                cp = tx.get("counterparty_name", "")
                inn = tx.get("counterparty_inn", "")
                purp = tx.get("purpose", "")
                bal = tx.get("closing_balance", "")
                cat = tx.get("category_id", "")
            else:
                d = tx.date.strftime("%Y-%m-%d") if hasattr(tx, "date") and isinstance(tx.date, (date, datetime)) else str(getattr(tx, "date", ""))
                amt = float(tx.amount) if hasattr(tx, "amount") and tx.amount is not None else ""
                dirn = getattr(tx, "direction", "")
                cp = getattr(tx, "counterparty_name", "")
                inn = getattr(tx, "counterparty_inn", "")
                purp = getattr(tx, "purpose", "")
                bal = float(tx.closing_balance) if hasattr(tx, "closing_balance") and tx.closing_balance is not None else ""
                cat = getattr(tx, "category_id", "")

            writer.writerow([
                d,
                amt,
                neutralize_formula(dirn),
                neutralize_formula(cp),
                neutralize_formula(inn),
                neutralize_formula(purp),
                bal,
                neutralize_formula(cat)
            ])
        return output.getvalue()

def case_income():
    from sqlalchemy import case
    return case(
        (TransactionRecord.direction == "income", TransactionRecord.amount),
        else_=0.0
    )

def case_expense():
    from sqlalchemy import case
    return case(
        (TransactionRecord.direction == "expense", TransactionRecord.amount),
        else_=0.0
    )
